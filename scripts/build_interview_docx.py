from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "ai_agent_interview_study_guide.md"
OUTPUT = ROOT / "docs" / "AI_Agent_实习面试项目学习文档.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_paragraph_border(paragraph, color: str = "D9E2EF") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "3")
        element.set(qn("w:color"), color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = styles["Title"]
    title.font.name = "Microsoft YaHei"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)
    title.paragraph_format.space_after = Pt(8)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Microsoft YaHei"
    subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    subtitle.font.size = Pt(11)
    subtitle.font.color.rgb = RGBColor(0x55, 0x65, 0x75)
    subtitle.paragraph_format.space_after = Pt(12)

    for style_name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header.paragraphs[0]
    header.text = "AI Agent 实习面试项目学习文档"
    header.style = styles["Header"]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x66, 0x72, 0x85)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    footer.style = styles["Footer"]
    for run in footer.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x66, 0x72, 0x85)


def add_title_block(document: Document) -> None:
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("AI Agent 实习面试项目学习文档")

    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Tarot Persona Agent 项目讲解、背诵稿与追问准备")

    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note_run = note.add_run("学习顺序：先背 2 分钟介绍，再理解 Agent 状态流，再理解 RAG 数据流，最后练习追问。")
    note_run.font.bold = True
    note_run.font.color.rgb = RGBColor(0x1F, 0x4D, 0x78)
    set_paragraph_shading(note, "E8EEF5")
    set_paragraph_border(note, "C7D7EA")


def add_markdown_text(paragraph, text: str) -> None:
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x9B, 0x1C, 0x1C)
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def add_code_block(document: Document, lines: list[str]) -> None:
    if not lines:
        return
    paragraph = document.add_paragraph()
    set_paragraph_shading(paragraph, "F4F6F9")
    set_paragraph_border(paragraph, "D9E2EF")
    text = "\n".join(lines)
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.05


def should_page_break_before(heading_text: str, level: int) -> bool:
    if level != 2:
        return False
    return heading_text.startswith(
        (
            "5. LangGraph",
            "6. RAG",
            "10. 高频",
            "13. 最后",
        )
    )


def build_docx() -> None:
    markdown = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()
    configure_document(document)
    add_title_block(document)

    in_code = False
    code_lines: list[str] = []
    in_blockquote = False

    for raw_line in markdown:
        line = raw_line.rstrip()

        if line.startswith("```"):
            if in_code:
                add_code_block(document, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue

        if in_code:
            code_lines.append(line)
            continue

        if not line.strip():
            in_blockquote = False
            continue

        if line.strip() == "---":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(6)
            run = paragraph.add_run(" " * 120)
            run.font.size = Pt(1)
            set_paragraph_border(paragraph, "E5E7EB")
            continue

        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        if heading:
            hashes, text = heading.groups()
            level = len(hashes)
            if level == 1:
                if text.strip() == "AI Agent 实习面试项目学习文档":
                    continue
                paragraph = document.add_paragraph(style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_markdown_text(paragraph, text)
            else:
                if should_page_break_before(text, level):
                    document.add_page_break()
                paragraph = document.add_paragraph(style=f"Heading {min(level - 1, 3)}")
                add_markdown_text(paragraph, text)
            continue

        if line.startswith(">"):
            quote_text = line.lstrip("> ").strip()
            if not quote_text:
                continue
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.2)
            paragraph.paragraph_format.right_indent = Inches(0.1)
            paragraph.paragraph_format.space_before = Pt(2 if in_blockquote else 6)
            paragraph.paragraph_format.space_after = Pt(4)
            set_paragraph_shading(paragraph, "FFF8E5")
            set_paragraph_border(paragraph, "E7C873")
            run = paragraph.add_run(quote_text)
            run.italic = True
            run.font.color.rgb = RGBColor(0x5F, 0x45, 0x00)
            in_blockquote = True
            continue

        unordered = re.match(r"^\s*-\s+(.*)$", line)
        if unordered:
            paragraph = document.add_paragraph(style="List Bullet")
            add_markdown_text(paragraph, unordered.group(1))
            continue

        ordered = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if ordered:
            paragraph = document.add_paragraph(style="List Number")
            add_markdown_text(paragraph, ordered.group(1))
            continue

        paragraph = document.add_paragraph()
        add_markdown_text(paragraph, line)

    if code_lines:
        add_code_block(document, code_lines)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_docx()
